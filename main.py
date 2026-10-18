import asyncio
import time as _time
import base64
import hashlib
import hmac
import json
import logging
import os
import secrets
import time
from contextlib import asynccontextmanager
from datetime import date, datetime, time as dtime, timedelta, timezone
from typing import Optional
from urllib.parse import parse_qsl
from zoneinfo import ZoneInfo

import httpx
from aiogram import Bot, Dispatcher, F
from aiogram.filters import CommandObject, CommandStart
from aiogram.types import (
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    KeyboardButton,
    MenuButtonWebApp,
    BufferedInputFile,
    CallbackQuery,
    Message,
    ReplyKeyboardMarkup,
    Update,
    WebAppInfo,
)
from fastapi import Depends, FastAPI, Header, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, Response
from pydantic import BaseModel

logging.basicConfig(level=logging.INFO)
log = logging.getLogger("barber")

# ---------------------------------------------------------------- config
BOT_TOKEN = os.environ["BOT_TOKEN"]
SUPABASE_URL = os.environ["SUPABASE_URL"].rstrip("/")
SUPABASE_KEY = os.environ["SUPABASE_KEY"]  # secret key (sb_secret_...) or service_role
OWNER_ID = int(os.environ["OWNER_ID"])  # barber's Telegram ID
WEBAPP_URL = os.environ.get("WEBAPP_URL", "")  # Netlify mini app URL
BASE_URL = (os.environ.get("BASE_URL") or os.environ.get("RENDER_EXTERNAL_URL", "")).rstrip("/")
WEBHOOK_SECRET = os.environ.get("WEBHOOK_SECRET") or hashlib.sha256(BOT_TOKEN.encode()).hexdigest()[:32]
TZ = ZoneInfo(os.environ.get("TIMEZONE", "Asia/Tashkent"))
SLOT_STEP = int(os.environ.get("SLOT_STEP_MIN", "15"))
MAX_ACTIVE_BOOKINGS = int(os.environ.get("MAX_ACTIVE_BOOKINGS", "3"))
CORS_ORIGINS = [o.strip() for o in os.environ.get("CORS_ORIGINS", "*").split(",")]

BOT_USERNAME = ""

# ---------------------------------------------------------------- bot messages (uz / ru)
MSG = {
    "uz": {
        "welcome": "Salom! Barberga yozilish uchun pastdagi tugmani bosing 👇",
        "btn_book": "✂️ Yozilish",
        "share_phone": "Barber siz bilan bog'lana olishi uchun telefon raqamingizni yuboring:",
        "btn_phone": "📱 Raqamni yuborish",
        "phone_saved": "Rahmat, raqam saqlandi ✅",
        "own_phone": "Iltimos, o'zingizning raqamingizni yuboring.",
        "admin_now": "Siz admin bo'ldingiz ✅",
        "invite_bad": "Bu havola eskirgan yoki ishlatilgan.",
        "reminder": "⏰ Eslatma: soat {t} da «{svc}» uchun yozilgansiz.\n📍 {branch}",
        "cancelled_by_staff": "❌ {when} dagi yozilishingiz bekor qilindi.",
        "staff_new": "🆕 Yangi yozilish\n{name} — {svc}\n📍 {branch}\n{when}{phone}",
        "staff_cancelled": "❌ Bekor qilindi: {name} — {svc}, {branch}, {when}",
    },
    "ru": {
        "welcome": "Здравствуйте! Нажмите кнопку ниже, чтобы записаться 👇",
        "btn_book": "✂️ Записаться",
        "share_phone": "Отправьте номер телефона, чтобы мастер мог с вами связаться:",
        "btn_phone": "📱 Отправить номер",
        "phone_saved": "Спасибо, номер сохранён ✅",
        "own_phone": "Пожалуйста, отправьте свой номер.",
        "admin_now": "Теперь вы администратор ✅",
        "invite_bad": "Ссылка устарела или уже использована.",
        "reminder": "⏰ Напоминание: в {t} у вас запись «{svc}».\n📍 {branch}",
        "cancelled_by_staff": "❌ Ваша запись на {when} отменена.",
        "staff_new": "🆕 Новая запись\n{name} — {svc}\n📍 {branch}\n{when}{phone}",
        "staff_cancelled": "❌ Отменено: {name} — {svc}, {branch}, {when}",
    },
}


MSG["uz"].update({
    "pending_client": "⏳ Usta tasdig'ini kutyapmiz\n\n✂️ Usta: {barber}\n📅 {when}\n\nXizmatlar ({n}):\n{svcs}",
    "confirmed": "✅ Yozuv tasdiqlandi!\n\n✂️ Usta: {barber}\n📅 {when}\n\nXizmatlar ({n}):\n{svcs}\n\n📍 Filial: {branch}\n🏠 {addr}{map}",
    "client_cancelled": "❌ Yozuv bekor qilindi\n\n{barber} ustaga {when} sanasidagi yozuv bekor qilindi.",
    "cancelled_by_staff": "❌ {when} dagi yozuvingiz usta tomonidan bekor qilindi.\nIltimos, boshqa vaqtga yozilishga urinib ko'ring.",
    "staff_new": "🆕 Yangi yozuv — tasdiqlash kerak\n{name} — {svc}\n✂️ {barber}\n📍 {branch}\n📅 {when}{phone}",
    "staff_cancelled": "❌ Mijoz o'zi bekor qildi: {name} — {svc}, {barber}, {when}",
    "btn_cancel": "❌ Bekor qilish", "btn_ok": "✅ Tasdiqlash", "btn_no": "❌ Rad etish",
})
MSG["ru"].update({
    "pending_client": "⏳ Ждём подтверждения мастера\n\n✂️ Мастер: {barber}\n📅 {when}\n\nУслуги ({n}):\n{svcs}",
    "confirmed": "✅ Запись подтверждена!\n\n✂️ Мастер: {barber}\n📅 {when}\n\nУслуги ({n}):\n{svcs}\n\n📍 Филиал: {branch}\n🏠 {addr}{map}",
    "client_cancelled": "❌ Запись отменена\n\nЗапись к мастеру {barber} на {when} отменена.",
    "cancelled_by_staff": "❌ Мастер отменил вашу запись на {when}.\nПожалуйста, попробуйте записаться на другое время.",
    "staff_new": "🆕 Новая запись — нужно подтвердить\n{name} — {svc}\n✂️ {barber}\n📍 {branch}\n📅 {when}{phone}",
    "staff_cancelled": "❌ Клиент сам отменил: {name} — {svc}, {barber}, {when}",
    "btn_cancel": "❌ Отменить", "btn_ok": "✅ Подтвердить", "btn_no": "❌ Отклонить",
})


MSG["uz"].update({
    "client_cancelled": "❌ Siz yozuvni bekor qildingiz\n\n✂️ Usta: {barber}\n📅 {when}\n\nBekor qilingan xizmatlar ({n}):\n{svcs}",
    "cancelled_by_staff": "❌ Yozuv usta tomonidan bekor qilindi\n\n✂️ Usta: {barber}\n📅 {when}\n\nBekor qilingan xizmatlar ({n}):\n{svcs}\n\nIltimos, boshqa vaqtga yozilishga urinib ko'ring.",
})
MSG["ru"].update({
    "client_cancelled": "❌ Вы отменили запись\n\n✂️ Мастер: {barber}\n📅 {when}\n\nОтменённые услуги ({n}):\n{svcs}",
    "cancelled_by_staff": "❌ Запись отменена мастером\n\n✂️ Мастер: {barber}\n📅 {when}\n\nОтменённые услуги ({n}):\n{svcs}\n\nПожалуйста, попробуйте записаться на другое время.",
})


MSG["uz"]["barber_now"] = "Siz usta sifatida ulandingiz ✅\nSizga yozilgan mijozlar haqida xabarlar shu yerga keladi."
MSG["ru"]["barber_now"] = "Вы подключены как мастер ✅\nУведомления о записях к вам будут приходить сюда."


async def edit_or_send(chat_id: int, msg_id, text: str):
    if msg_id:
        try:
            await bot.edit_message_text(text, chat_id=chat_id, message_id=msg_id)
            return
        except Exception:
            log.warning("edit failed for %s", chat_id)
    await notify(chat_id, text)


def loc(row: dict, key: str, lang: Optional[str]) -> str:
    return ((row.get(key + "_ru") if lang == "ru" else None) or row.get(key) or "")


def tr(lang: Optional[str], key: str, **kw) -> str:
    return (MSG.get(lang or "uz") or MSG["uz"])[key].format(**kw)


# ---------------------------------------------------------------- database (Supabase REST)
class DBError(Exception):
    def __init__(self, status: int, body: str):
        super().__init__(f"{status}: {body}")
        self.status = status
        self.body = body


_cache: dict = {}
_bg: set = set()


async def cached(key, ttl, tables, fn):
    hit = _cache.get(key)
    now = _time.monotonic()
    if hit and hit[0] > now:
        return hit[1]
    val = await fn()
    _cache[key] = (now + ttl, val, frozenset(tables))
    return val


def invalidate(table: str):
    for k in [k for k, v in _cache.items() if table in v[2]]:
        _cache.pop(k, None)


def bg(coro):
    t = asyncio.create_task(coro)
    _bg.add(t)
    t.add_done_callback(_bg.discard)


class DB:
    def __init__(self, url: str, key: str):
        headers = {"apikey": key, "Content-Type": "application/json"}
        if not key.startswith("sb_"):  # legacy JWT keys also need Authorization
            headers["Authorization"] = f"Bearer {key}"
        self.c = httpx.AsyncClient(base_url=f"{url}/rest/v1", headers=headers, timeout=20,
                                   limits=httpx.Limits(max_keepalive_connections=20, keepalive_expiry=60))

    async def _req(self, method, table, params=None, body=None):
        r = await self.c.request(
            method, f"/{table}", params=params, json=body,
            headers={"Prefer": "return=representation"},
        )
        if r.status_code >= 400:
            raise DBError(r.status_code, r.text)
        if method != "GET":
            invalidate(table)
        return r.json() if r.text else []

    async def select(self, table, params=None):
        return await self._req("GET", table, params=params)

    async def insert(self, table, data):
        return await self._req("POST", table, body=data)

    async def update(self, table, filters, data):
        return await self._req("PATCH", table, params=filters, body=data)

    async def delete(self, table, filters):
        return await self._req("DELETE", table, params=filters)

    async def close(self):
        await self.c.aclose()


db = DB(SUPABASE_URL, SUPABASE_KEY)
bot = Bot(BOT_TOKEN)
dp = Dispatcher()


# ---------------------------------------------------------------- helpers
def pdt(s: str) -> datetime:
    return datetime.fromisoformat(s.replace("Z", "+00:00"))


def now_utc_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def fmt_dt(dt: datetime) -> str:
    return dt.astimezone(TZ).strftime("%d.%m %H:%M")


async def notify(chat_id: int, text: str, markup=None):
    try:
        return await bot.send_message(chat_id, text, reply_markup=markup)
    except Exception:
        log.warning("notify failed for %s", chat_id)


def kb(*rows):
    return InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text=t, callback_data=d) for t, d in r] for r in rows])


async def staff_users() -> list[dict]:
    return await cached("staff", 60, {"users"}, lambda: db.select("users", {"role": "in.(owner,admin)", "select": "telegram_id,lang,role"}))


async def get_lang(tid: int):
    rows = await cached(f"user:{tid}", 60, {"users"}, lambda: db.select("users", {"telegram_id": f"eq.{tid}"}))
    return rows[0].get("lang") if rows else None


async def staff_recipients(barber_id=None) -> list[dict]:
    """barber_id berilsa: faqat egasi (owner) + o'sha usta. Berilmasa: barcha owner/adminlar."""
    staff = await staff_users()
    if barber_id is None:
        return staff
    out = [s for s in staff if s.get("role") == "owner"]
    if not out:
        out = [{"telegram_id": OWNER_ID, "lang": None}]
    rows = await cached(f"barber_tg:{barber_id}", 60, {"barbers"},
                        lambda: db.select("barbers", {"id": f"eq.{barber_id}", "select": "telegram_id"}))
    tid = rows[0].get("telegram_id") if rows else None
    if tid and not any(s["telegram_id"] == tid for s in out):
        out.append({"telegram_id": tid, "lang": await get_lang(tid)})
    return out


def _kw_lang(kw: dict, lg) -> dict:
    return {k: (((v[1] or v[0]) if lg == "ru" else v[0]) if isinstance(v, tuple) else v) for k, v in kw.items()}


async def notify_staff(key: str, markup_fn=None, barber_id=None, **kw) -> dict:
    """Egasiga va (barber_id berilsa) o'sha ustaga yuboradi. Qaytaradi: {chat_id(str): message_id}"""
    sent = {}
    for s in await staff_recipients(barber_id):
        lg = s.get("lang")
        m = await notify(s["telegram_id"], tr(lg, key, **_kw_lang(kw, lg)), markup_fn(lg) if markup_fn else None)
        if m:
            sent[str(s["telegram_id"])] = m.message_id
    return sent


async def edit_staff_msgs(b: dict, key: str, **kw) -> bool:
    """Bron uchun adminlarga yuborilgan xabarlarni tahrirlaydi (tugmalar ham o'chadi)."""
    msgs = b.get("staff_msgs") or {}
    if not msgs:
        return False
    langs = {str(s["telegram_id"]): s.get("lang") for s in await staff_users()}
    for cid, mid in msgs.items():
        lg = langs[cid] if cid in langs else await get_lang(int(cid))
        await edit_or_send(int(cid), mid, tr(lg, key, **_kw_lang(kw, lg)))
    return True


async def ensure_user(tid: int, name: str, lang_hint: Optional[str] = None) -> dict:
    rows = await cached(f"user:{tid}", 60, {"users"}, lambda: db.select("users", {"telegram_id": f"eq.{tid}"}))
    if rows:
        u = rows[0]
    else:
        lang = "ru" if (lang_hint or "").startswith("ru") else "uz"
        u = (await db.insert("users", {"telegram_id": tid, "name": name, "role": "client", "lang": lang}))[0]
    if tid == OWNER_ID and u["role"] != "owner":
        u = (await db.update("users", {"telegram_id": f"eq.{tid}"}, {"role": "owner"}))[0]
    return u


def img_url(path: str, ver: int) -> Optional[str]:
    return f"{path}?v={ver}" if ver else None


def branch_out(r: dict) -> dict:
    return {
        "id": r["id"], "name": r["name"], "name_ru": r.get("name_ru"), "address": r.get("address"), "address_ru": r.get("address_ru"),
        "lat": r.get("lat"), "lng": r.get("lng"), "active": r.get("active", True),
        "photo_url": img_url(f"/api/img/branch/{r['id']}", r.get("photo_ver") or 0),
    }


BRANCH_COLS = "id,name,name_ru,address,address_ru,lat,lng,active,photo_ver"
SHOP_COLS = "name,name_ru,logo_ver,desc_uz,desc_ru,address,address_ru,phone,lat,lng"


def shop_out(r: dict) -> dict:
    return {
        "name": r["name"], "name_ru": r.get("name_ru"), "address_ru": r.get("address_ru"), "desc_uz": r.get("desc_uz"), "desc_ru": r.get("desc_ru"),
        "address": r.get("address"), "phone": r.get("phone"),
        "lat": r.get("lat"), "lng": r.get("lng"),
        "logo_url": img_url("/api/img/logo", r.get("logo_ver") or 0),
    }


BARBER_COLS = "id,branch_id,name,name_ru,active,photo_ver"
BARBER_COLS_ADM = BARBER_COLS + ",telegram_id"


def barber_out(r: dict, min_price=None) -> dict:
    return {"id": r["id"], "branch_id": r["branch_id"], "name": r["name"], "name_ru": r.get("name_ru"), "active": r.get("active", True),
            "min_price": min_price, "photo_url": img_url(f"/api/img/barber/{r['id']}", r.get("photo_ver") or 0)}


def barber_admin_out(r: dict) -> dict:
    return {**barber_out(r), "linked": bool(r.get("telegram_id"))}


def check_coords(data: dict):
    for k, lim in (("lat", 90), ("lng", 180)):
        v = data.get(k)
        if v is not None and not -lim <= v <= lim:
            raise HTTPException(400, "Koordinata noto'g'ri")
    if (data.get("lat") is None) != (data.get("lng") is None) and ("lat" in data and "lng" in data):
        raise HTTPException(400, "Koordinata noto'g'ri")


def check_image(data: str):
    if not data.startswith(("data:image/jpeg;base64,", "data:image/png;base64,", "data:image/webp;base64,")):
        raise HTTPException(400, "Rasm formati noto'g'ri")
    if len(data) > 900_000:
        raise HTTPException(400, "Rasm juda katta")
    try:
        base64.b64decode(data.split(",", 1)[1], validate=True)
    except Exception:
        raise HTTPException(400, "Rasm buzilgan")


def image_response(data_url: Optional[str]) -> Response:
    if not data_url:
        raise HTTPException(404)
    head, b64 = data_url.split(",", 1)
    mime = head[5:].split(";")[0]
    return Response(
        base64.b64decode(b64), media_type=mime,
        headers={"Cache-Control": "public, max-age=604800"},
    )


# ---------------------------------------------------------------- Telegram initData auth
def verify_init_data(init_data: str) -> dict:
    if not init_data:
        raise HTTPException(401, "auth_required")
    data = dict(parse_qsl(init_data, keep_blank_values=True))
    recv_hash = data.pop("hash", None)
    if not recv_hash:
        raise HTTPException(401, "auth_invalid")
    check = "\n".join(f"{k}={v}" for k, v in sorted(data.items()))
    secret = hmac.new(b"WebAppData", BOT_TOKEN.encode(), hashlib.sha256).digest()
    calc = hmac.new(secret, check.encode(), hashlib.sha256).hexdigest()
    if not hmac.compare_digest(calc, recv_hash):
        raise HTTPException(401, "auth_invalid")
    if time.time() - int(data.get("auth_date", 0)) > 86400:
        raise HTTPException(401, "auth_expired")
    return json.loads(data["user"])


async def get_user(x_init_data: str = Header(default="")) -> dict:
    tg = verify_init_data(x_init_data)
    name = f"{tg.get('first_name', '')} {tg.get('last_name', '')}".strip()
    return await ensure_user(tg["id"], name, tg.get("language_code"))


async def admin_user(u: dict = Depends(get_user)) -> dict:
    if u["role"] not in ("owner", "admin"):
        raise HTTPException(403, "Ruxsat yo'q")
    return u


async def barber_ids_of(tid: int) -> list:
    rows = await cached(f"barber_ids:{tid}", 60, {"barbers"},
                        lambda: db.select("barbers", {"telegram_id": f"eq.{tid}", "select": "id"}))
    return [r["id"] for r in rows]


async def staff_user(u: dict = Depends(get_user)) -> dict:
    """Owner/admin (scope=None, hammasi) yoki Telegramga ulangan usta (scope=[o'z barber id'lari])."""
    if u["role"] in ("owner", "admin"):
        return {**u, "scope": None}
    ids = await barber_ids_of(u["telegram_id"])
    if not ids:
        raise HTTPException(403, "Ruxsat yo'q")
    return {**u, "scope": ids}


def in_list(ids: list) -> str:
    return "in.(" + ",".join(str(int(i)) for i in ids) + ")"


def need_own_barber(a: dict, barber_id) -> None:
    if a["scope"] is not None and barber_id not in a["scope"]:
        raise HTTPException(403, "Ruxsat yo'q")


async def owner_user(u: dict = Depends(get_user)) -> dict:
    if u["role"] != "owner":
        raise HTTPException(403, "Faqat egasi uchun")
    return u


# ---------------------------------------------------------------- slots
async def get_service(sid: int, active_only: bool = True) -> dict:
    rows = await db.select("services", {"id": f"eq.{sid}"})
    if not rows or (active_only and not rows[0]["active"]):
        raise HTTPException(404, "not_found")
    return rows[0]


async def get_branch(bid: int, active_only: bool = True) -> dict:
    rows = await cached(f"branch:{bid}", 60, {"branches"}, lambda: db.select("branches", {"id": f"eq.{bid}", "select": BRANCH_COLS}))
    if not rows or (active_only and not rows[0]["active"]):
        raise HTTPException(404, "not_found")
    return rows[0]


async def barber_parts(barber_id: int, ids_csv: str):
    try:
        ids = sorted({int(x) for x in ids_csv.split(",") if x})
    except ValueError:
        raise HTTPException(404, "not_found")
    if not ids:
        raise HTTPException(404, "not_found")
    bar, svcs = await asyncio.gather(
        cached(f"bar:{barber_id}", 60, {"barbers"}, lambda: db.select("barbers", {"id": f"eq.{barber_id}", "active": "eq.true"})),
        cached(f"svc:{barber_id}:{ids}", 60, {"services"}, lambda: db.select("services", {"id": f"in.({','.join(map(str, ids))})", "active": "eq.true", "barber_id": f"eq.{barber_id}"})),
    )
    if not bar or len(svcs) != len(ids):
        raise HTTPException(404, "not_found")
    return bar[0], svcs


async def load_busy(barber_id: int, d: date, lo: datetime, fresh: bool):
    def q():
        return db.select("bookings", [("barber_id", f"eq.{barber_id}"), ("status", "in.(pending,booked)"),
                                      ("start_at", f"lt.{(lo + timedelta(days=1)).isoformat()}"), ("end_at", f"gt.{lo.isoformat()}")])
    return await q() if fresh else await cached(f"busy:{barber_id}:{d}", 30, {"bookings"}, q)


async def compute_slots(d: date, dur_min: int, branch_id: int, barber_id: int, fresh: bool = False) -> dict:
    empty = {"open": None, "close": None, "slots": []}
    lo = datetime.combine(d, dtime.min, TZ)
    off, wh, booked = await asyncio.gather(
        cached(f"dayoff:{barber_id}:{d}", 30, {"days_off"}, lambda: db.select("days_off", {"barber_id": f"eq.{barber_id}", "day": f"eq.{d.isoformat()}"})),
        cached(f"wh:{branch_id}:{d.weekday()}", 60, {"work_hours"}, lambda: db.select("work_hours", {"branch_id": f"eq.{branch_id}", "weekday": f"eq.{d.weekday()}"})),
        load_busy(barber_id, d, lo, fresh),
    )
    if off or not wh or wh[0]["is_off"] or not wh[0]["start_time"] or not wh[0]["end_time"]:
        return empty
    st = dtime.fromisoformat(wh[0]["start_time"])
    en = dtime.fromisoformat(wh[0]["end_time"])
    day_start = datetime.combine(d, st, TZ)
    day_end = datetime.combine(d, en, TZ)
    dur = timedelta(minutes=dur_min)
    busy = [(pdt(b["start_at"]), pdt(b["end_at"])) for b in booked]
    earliest = datetime.now(TZ) + timedelta(minutes=30)
    out, cur = [], day_start
    while cur + dur <= day_end:
        if cur >= earliest and all(not (cur < be and cur + dur > bs) for bs, be in busy):
            out.append(cur.isoformat())
        cur += timedelta(minutes=SLOT_STEP)
    return {"open": st.strftime("%H:%M"), "close": en.strftime("%H:%M"), "slots": out, "step": SLOT_STEP}


# ---------------------------------------------------------------- models
class BookingIn(BaseModel):
    barber_id: int
    service_ids: list[int]
    start_at: str
    phone: Optional[str] = None


class BarberIn(BaseModel):
    branch_id: int
    name: str
    name_ru: Optional[str] = None


class BarberPatch(BaseModel):
    branch_id: Optional[int] = None
    name: Optional[str] = None
    name_ru: Optional[str] = None
    active: Optional[bool] = None


class DayOffIn(BaseModel):
    barber_id: int
    day: str
    off: bool = True


class LangIn(BaseModel):
    lang: str


class StatusIn(BaseModel):
    status: str


class ServiceIn(BaseModel):
    barber_id: Optional[int] = None
    name: str
    name_ru: Optional[str] = None
    price: int
    duration_min: int


class ServicePatch(BaseModel):
    name: Optional[str] = None
    name_ru: Optional[str] = None
    price: Optional[int] = None
    duration_min: Optional[int] = None
    active: Optional[bool] = None


class HoursIn(BaseModel):
    start_time: Optional[str] = None
    end_time: Optional[str] = None
    is_off: bool = False


class BranchIn(BaseModel):
    name: str
    name_ru: Optional[str] = None
    address_ru: Optional[str] = None
    address: Optional[str] = None
    lat: Optional[float] = None
    lng: Optional[float] = None


class BranchPatch(BaseModel):
    name: Optional[str] = None
    name_ru: Optional[str] = None
    address_ru: Optional[str] = None
    address: Optional[str] = None
    lat: Optional[float] = None
    lng: Optional[float] = None
    active: Optional[bool] = None


class ShopPatch(BaseModel):
    name: Optional[str] = None
    name_ru: Optional[str] = None
    address_ru: Optional[str] = None
    desc_uz: Optional[str] = None
    desc_ru: Optional[str] = None
    address: Optional[str] = None
    phone: Optional[str] = None
    lat: Optional[float] = None
    lng: Optional[float] = None


class ImageIn(BaseModel):
    data: str


# ---------------------------------------------------------------- Telegram bot handlers
def main_keyboard(lang: str) -> Optional[InlineKeyboardMarkup]:
    if not WEBAPP_URL:
        return None
    return InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text=tr(lang, "btn_book"), web_app=WebAppInfo(url=WEBAPP_URL))
    ]])


@dp.message(CommandStart())
async def cmd_start(m: Message, command: CommandObject):
    u = await ensure_user(m.from_user.id, m.from_user.full_name, m.from_user.language_code)

    arg = command.args or ""
    if arg.startswith("brb_"):
        inv = await db.update("invites", {"code": f"eq.{arg[4:]}", "used": "eq.false", "barber_id": "not.is.null"}, {"used": True})
        if inv:
            await db.update("barbers", {"id": f"eq.{inv[0]['barber_id']}"}, {"telegram_id": u["telegram_id"]})
            await m.answer(tr(u.get("lang"), "barber_now"))
        else:
            await m.answer(tr(u.get("lang"), "invite_bad"))
    elif arg.startswith("adm_") and u["role"] == "client":
        code = arg[4:]
        used = await db.update("invites", {"code": f"eq.{code}", "used": "eq.false"}, {"used": True})
        if used:
            u = (await db.update("users", {"telegram_id": f"eq.{u['telegram_id']}"}, {"role": "admin"}))[0]
            await m.answer(tr(u.get("lang"), "admin_now"))
        else:
            await m.answer(tr(u.get("lang"), "invite_bad"))

    lang = u.get("lang") or "uz"
    await m.answer(tr(lang, "welcome"), reply_markup=main_keyboard(lang))
    if not u.get("phone"):
        kb = ReplyKeyboardMarkup(
            keyboard=[[KeyboardButton(text=tr(lang, "btn_phone"), request_contact=True)]],
            resize_keyboard=True, one_time_keyboard=True,
        )
        await m.answer(tr(lang, "share_phone"), reply_markup=kb)


@dp.message(F.contact)
async def on_contact(m: Message):
    u = await ensure_user(m.from_user.id, m.from_user.full_name, m.from_user.language_code)
    lang = u.get("lang") or "uz"
    c = m.contact
    if c.user_id != m.from_user.id:
        await m.answer(tr(lang, "own_phone"))
        return
    await db.update("users", {"telegram_id": f"eq.{m.from_user.id}"}, {"phone": c.phone_number})
    await m.answer(tr(lang, "phone_saved"), reply_markup=main_keyboard(lang))


# ---------------------------------------------------------------- booking status flow
def map_link(b: dict) -> str:
    br = b.get("branches") or {}
    return f"\n🗺 https://www.google.com/maps/search/?api=1&query={br['lat']},{br['lng']}" if br.get("lat") is not None and br.get("lng") is not None else ""


async def transition(bid: int, new: str, actor: dict) -> dict:
    rows = await db.select("bookings", [("id", f"eq.{bid}"),
                           ("select", "*,barbers(name,name_ru,telegram_id),branches(name,name_ru,address,address_ru,lat,lng),users(name,lang)")])
    if not rows:
        raise HTTPException(404, "not_found")
    b = rows[0]
    own_barber = (b.get("barbers") or {}).get("telegram_id")
    staff = actor["role"] in ("owner", "admin") or (own_barber is not None and own_barber == actor["telegram_id"])
    if not staff and b["client_id"] != actor["telegram_id"]:
        raise HTTPException(403, "forbidden")
    if new == "booked" and (not staff or b["status"] != "pending"):
        raise HTTPException(400, "cannot_cancel")
    if new == "cancelled":
        if b["status"] not in ("pending", "booked"):
            raise HTTPException(400, "cannot_cancel")
        if not staff and pdt(b["start_at"]) < datetime.now(timezone.utc):
            raise HTTPException(400, "too_late")
    out = (await db.update("bookings", {"id": f"eq.{bid}"}, {"status": new}))[0]
    lang, when = (b["users"] or {}).get("lang"), fmt_dt(pdt(b["start_at"]))
    barber = loc(b.get("barbers") or {}, "name", lang)
    names = ((b.get("svc_names_ru") if lang == "ru" else None) or b.get("svc_names") or "").split(", ")
    svcs = "\n".join("• " + n for n in names)
    if new == "booked":
        await edit_or_send(b["client_id"], b.get("client_msg_id"), tr(lang, "confirmed", barber=barber, when=when, n=len(names), svcs=svcs,
                           branch=loc(b["branches"], "name", lang), addr=loc(b["branches"], "address", lang), map=map_link(b)))
    else:
        cancelled = "\n".join(f"• {n} - {when}" for n in names)
        if b["client_id"] == actor["telegram_id"]:
            await edit_or_send(b["client_id"], b.get("client_msg_id"), tr(lang, "client_cancelled", barber=barber, when=when, n=len(names), svcs=cancelled))
            kw = dict(name=b["users"]["name"], svc=(b.get("svc_names") or "", b.get("svc_names_ru") or b.get("svc_names") or ""),
                      barber=(loc(b.get("barbers") or {}, "name", "uz"), loc(b.get("barbers") or {}, "name", "ru")), when=when)
            # tasdiqlanmagan bo'lsa -- eski xabarni almashtiramiz (yangi xabar yubormaymiz);
            # tasdiqlangan bo'lsa -- admin bilishi uchun yangi xabar yuboramiz
            if not (b["status"] == "pending" and await edit_staff_msgs(b, "staff_cancelled", **kw)):
                await notify_staff("staff_cancelled", barber_id=b["barber_id"], **kw)
        else:
            await edit_or_send(b["client_id"], b.get("client_msg_id"), tr(lang, "cancelled_by_staff", barber=barber, when=when, n=len(names), svcs=cancelled))
    return out


@dp.callback_query(F.data.regexp(r"^(c|ok|no):\d+$"))
async def on_booking_cb(cq: CallbackQuery):
    kind, bid = cq.data.split(":")
    u = await ensure_user(cq.from_user.id, cq.from_user.full_name, cq.from_user.language_code)
    try:
        out = await transition(int(bid), "booked" if kind == "ok" else "cancelled", u)
    except HTTPException as e:
        await cq.answer(str(e.detail), show_alert=True)
        return
    await cq.answer("✅")
    try:
        await cq.message.edit_reply_markup(reply_markup=None)
    except Exception:
        pass
    if kind in ("ok", "no"):  # ikkinchi qabul qiluvchida (egasi/usta) tugmalar qolib ketmasin
        for cid, mid in (out.get("staff_msgs") or {}).items():
            if int(cid) != cq.message.chat.id:
                try:
                    await bot.edit_message_reply_markup(chat_id=int(cid), message_id=mid, reply_markup=None)
                except Exception:
                    pass


# ---------------------------------------------------------------- reminders
async def reminder_loop():
    while True:
        try:
            now = datetime.now(timezone.utc).replace(microsecond=0)
            rows = await db.select("bookings", [
                ("status", "eq.booked"),
                ("reminded", "eq.false"),
                ("start_at", f"gt.{now.isoformat()}"),
                ("start_at", f"lte.{(now + timedelta(hours=2)).isoformat()}"),
                ("select", "*,services(name),branches(name,name_ru),users(lang)"),
            ])
            for r in rows:
                await notify(r["client_id"], tr(
                    (r.get("users") or {}).get("lang"), "reminder",
                    t=pdt(r["start_at"]).astimezone(TZ).strftime("%H:%M"),
                    svc=((r.get("svc_names_ru") if (r.get("users") or {}).get("lang") == "ru" else None) or r.get("svc_names") or r["services"]["name"]),
                    branch=loc(r.get("branches") or {}, "name", (r.get("users") or {}).get("lang")),
                ))
                await db.update("bookings", {"id": f"eq.{r['id']}"}, {"reminded": True})
        except Exception:
            log.exception("reminder loop error")
        await asyncio.sleep(60)


# ---------------------------------------------------------------- app
@asynccontextmanager
async def lifespan(app: FastAPI):
    global BOT_USERNAME
    BOT_USERNAME = (await bot.get_me()).username
    if BASE_URL:
        await bot.set_webhook(f"{BASE_URL}/webhook/{WEBHOOK_SECRET}", secret_token=WEBHOOK_SECRET)
    if WEBAPP_URL:
        await bot.set_chat_menu_button(
            menu_button=MenuButtonWebApp(text="✂️ Yozilish", web_app=WebAppInfo(url=WEBAPP_URL))
        )
    task = asyncio.create_task(reminder_loop())
    yield
    task.cancel()
    await bot.session.close()
    await db.close()


app = FastAPI(lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=CORS_ORIGINS, allow_methods=["*"], allow_headers=["*"])


@app.exception_handler(DBError)
async def db_error_handler(request: Request, exc: DBError):
    log.error("DB error %s", exc)
    return JSONResponse({"detail": "server_error"}, status_code=500)


@app.api_route("/health", methods=["GET", "HEAD"])
async def health():
    return {"ok": True}


@app.post("/webhook/{secret}")
async def webhook(secret: str, request: Request):
    if secret != WEBHOOK_SECRET or request.headers.get("X-Telegram-Bot-Api-Secret-Token") != WEBHOOK_SECRET:
        raise HTTPException(403)
    update = Update.model_validate(await request.json(), context={"bot": bot})
    await dp.feed_update(bot, update)
    return {"ok": True}


# ---- public images (no auth: they are shop photos)
@app.get("/api/img/logo")
async def img_logo():
    rows = await db.select("shop", {"id": "eq.1", "select": "logo_data"})
    return image_response(rows[0]["logo_data"] if rows else None)


@app.get("/api/img/branch/{bid}")
async def img_branch(bid: int):
    rows = await db.select("branches", {"id": f"eq.{bid}", "select": "photo_data"})
    return image_response(rows[0]["photo_data"] if rows else None)


# ---------------------------------------------------------------- client API
async def public_data():
    shop, branches, services, barbers = await asyncio.gather(
        db.select("shop", {"id": "eq.1", "select": SHOP_COLS}),
        db.select("branches", {"active": "eq.true", "order": "id.asc", "select": BRANCH_COLS}),
        db.select("services", {"active": "eq.true", "order": "id.asc"}),
        db.select("barbers", {"active": "eq.true", "order": "id.asc", "select": BARBER_COLS}),
    )
    mins = {}
    for sv in services:
        if sv.get("barber_id"):
            mins[sv["barber_id"]] = min(mins.get(sv["barber_id"], sv["price"]), sv["price"])
    return {
        "shop": shop_out(shop[0]) if shop else shop_out({"name": "Barbershop"}),
        "branches": [branch_out(b) for b in branches],
        "barbers": [barber_out(b, mins.get(b["id"])) for b in barbers],
        "services": services,
    }


@app.get("/api/bootstrap")
async def bootstrap(u: dict = Depends(get_user)):
    pub = await cached("public", 120, {"shop", "branches", "services", "barbers"}, public_data)
    return {"me": {"id": u["telegram_id"], "name": u["name"], "role": u["role"],
                   "barber": bool(await barber_ids_of(u["telegram_id"])),
                   "phone": u.get("phone"), "lang": u.get("lang") or "uz"}, **pub}


@app.post("/api/me/lang")
async def set_lang(body: LangIn, u: dict = Depends(get_user)):
    if body.lang not in ("uz", "ru"):
        raise HTTPException(400, "bad_lang")
    await db.update("users", {"telegram_id": f"eq.{u['telegram_id']}"}, {"lang": body.lang})
    return {"ok": True}


@app.get("/api/services")
async def services(u: dict = Depends(get_user)):
    return await db.select("services", {"active": "eq.true", "order": "id.asc"})


@app.get("/api/barber-off")
async def barber_off(barber_id: int, u: dict = Depends(get_user)):
    async def load():
        bar = await db.select("barbers", {"id": f"eq.{barber_id}"})
        if not bar:
            raise HTTPException(404, "not_found")
        wh, days = await asyncio.gather(
            db.select("work_hours", {"branch_id": f"eq.{bar[0]['branch_id']}", "is_off": "eq.true"}),
            db.select("days_off", {"barber_id": f"eq.{barber_id}", "day": f"gte.{datetime.now(TZ).date().isoformat()}"}))
        return {"weekdays": [w["weekday"] for w in wh], "days": [d["day"] for d in days]}
    return await cached(f"off:{barber_id}:{datetime.now(TZ).date()}", 30, {"barbers", "work_hours", "days_off"}, load)


@app.get("/api/slots")
async def slots(date_: str, barber_id: int, service_ids: str, u: dict = Depends(get_user)):
    try:
        d = date.fromisoformat(date_)
    except ValueError:
        raise HTTPException(400, "bad_date")
    today = datetime.now(TZ).date()
    if d < today or d > today + timedelta(days=30):
        return {"open": None, "close": None, "slots": []}
    bar, svcs = await barber_parts(barber_id, service_ids)
    return await compute_slots(d, sum(x["duration_min"] for x in svcs), bar["branch_id"], bar["id"])


@app.post("/api/bookings")
async def create_booking(b: BookingIn, u: dict = Depends(get_user)):
    bar, svcs = await barber_parts(b.barber_id, ",".join(map(str, b.service_ids)))
    br = await get_branch(bar["branch_id"])
    try:
        start = datetime.fromisoformat(b.start_at)
    except ValueError:
        raise HTTPException(400, "bad_time")
    if start.tzinfo is None:
        start = start.replace(tzinfo=TZ)
    start = start.astimezone(TZ)
    dur = sum(x["duration_min"] for x in svcs)
    slot_res, active = await asyncio.gather(
        compute_slots(start.date(), dur, br["id"], bar["id"], fresh=True),
        db.select("bookings", [("client_id", f"eq.{u['telegram_id']}"), ("status", "in.(pending,booked)"),
                               ("start_at", f"gt.{now_utc_iso()}"), ("select", "id")]))
    free = [datetime.fromisoformat(x) for x in slot_res["slots"]]
    if start not in free:
        raise HTTPException(409, "slot_unavailable")
    if len(active) >= MAX_ACTIVE_BOOKINGS:
        raise HTTPException(400, "too_many")
    phone = (b.phone or "").strip()[:32] or u.get("phone")
    if phone and phone != u.get("phone"):
        bg(db.update("users", {"telegram_id": f"eq.{u['telegram_id']}"}, {"phone": phone}))
    names = ", ".join(x["name"] for x in svcs)
    names_ru = ", ".join(loc(x, "name", "ru") for x in svcs)
    try:
        rows = await db.insert("bookings", {
            "client_id": u["telegram_id"], "service_id": svcs[0]["id"], "service_ids": [x["id"] for x in svcs],
            "barber_id": bar["id"], "branch_id": br["id"], "status": "pending", "svc_names": names, "svc_names_ru": names_ru,
            "total_price": sum(x["price"] for x in svcs), "total_min": dur,
            "start_at": start.isoformat(), "end_at": (start + timedelta(minutes=dur)).isoformat(),
            "reminded": (start - datetime.now(TZ)) <= timedelta(hours=2),
        })
    except DBError as e:
        if "23P01" in e.body:
            raise HTTPException(409, "slot_taken")
        raise
    bk = rows[0]
    when, lang = fmt_dt(start), u.get("lang")
    async def send_pending():
        m = await notify(u["telegram_id"], tr(lang, "pending_client", barber=loc(bar, "name", lang), when=when, n=len(svcs),
              svcs="\n".join("• " + loc(x, "name", lang) for x in svcs)), kb([(tr(lang, "btn_cancel"), f"c:{bk['id']}")]))
        if m:
            await db.update("bookings", {"id": f"eq.{bk['id']}"}, {"client_msg_id": m.message_id})
    bg(send_pending())
    async def send_staff():
        sent = await notify_staff("staff_new", lambda l: kb([(tr(l, "btn_ok"), f"ok:{bk['id']}"), (tr(l, "btn_no"), f"no:{bk['id']}")]),
                       barber_id=bar["id"],
                       name=u["name"], svc=(names, names_ru), barber=(bar["name"], loc(bar, "name", "ru")), branch=(br["name"], loc(br, "name", "ru")), when=when,
                       phone=(f"\n📞 {phone}" if phone else ""))
        if sent:
            try:
                await db.update("bookings", {"id": f"eq.{bk['id']}"}, {"staff_msgs": sent})
            except Exception:
                log.warning("staff_msgs saqlanmadi (bookings.staff_msgs ustuni bormi?)")
    bg(send_staff())
    return bk


@app.get("/api/my-bookings")
async def my_bookings(scope: str = "active", u: dict = Depends(get_user)):
    params = [("client_id", f"eq.{u['telegram_id']}"), ("select", "*,barbers(name,name_ru),branches(name,name_ru,address,address_ru,lat,lng)")]
    now_iso = now_utc_iso()
    if scope == "history":
        params += [("or", f"(status.in.(cancelled,done,no_show),end_at.lte.{now_iso})"),
                   ("order", "start_at.desc"), ("limit", "50")]
    else:
        params += [("status", "in.(pending,booked)"), ("end_at", f"gt.{now_iso}"), ("order", "start_at.asc")]
    return await db.select("bookings", params)


@app.post("/api/bookings/{bid}/cancel")
async def cancel_booking(bid: int, u: dict = Depends(get_user)):
    await transition(bid, "cancelled", u)
    return {"ok": True}


# ---- barbers (public photo + admin)
@app.get("/api/img/barber/{bid}")
async def img_barber(bid: int):
    rows = await db.select("barbers", {"id": f"eq.{bid}", "select": "photo_data"})
    return image_response(rows[0]["photo_data"] if rows else None)


# ---------------------------------------------------------------- admin API
@app.get("/api/admin/bookings")
async def admin_bookings(date_: str, branch_id: int = 0, barber_id: int = 0, a: dict = Depends(staff_user)):
    try:
        d = date.fromisoformat(date_)
    except ValueError:
        raise HTTPException(400, "Sana noto'g'ri")
    s = datetime.combine(d, dtime(0, 0), TZ)
    e = s + timedelta(days=1)
    params = [
        ("start_at", f"gte.{s.isoformat()}"), ("start_at", f"lt.{e.isoformat()}"),
        ("order", "start_at.asc"), ("select", "*,barbers(name,name_ru),users(name,phone),branches(name,name_ru)"),
    ]
    if branch_id:
        params.append(("branch_id", f"eq.{branch_id}"))
    if barber_id:
        params.append(("barber_id", f"eq.{barber_id}"))
    if a["scope"] is not None:
        params.append(("barber_id", in_list(a["scope"])))
    return await db.select("bookings", params)


@app.post("/api/admin/bookings/{bid}/status")
async def admin_set_status(bid: int, body: StatusIn, a: dict = Depends(staff_user)):
    if a["scope"] is not None:
        cur = await db.select("bookings", {"id": f"eq.{bid}", "select": "barber_id"})
        if not cur:
            raise HTTPException(404, "Topilmadi")
        need_own_barber(a, cur[0]["barber_id"])
    if body.status in ("booked", "cancelled"):
        out = await transition(bid, body.status, a)
        for cid, mid in (out.get("staff_msgs") or {}).items():  # botdagi xabarlarda tugmalar qolib ketmasin
            try:
                await bot.edit_message_reply_markup(chat_id=int(cid), message_id=mid, reply_markup=None)
            except Exception:
                pass
        return out
    if body.status not in ("done", "no_show"):
        raise HTTPException(400, "Noto'g'ri holat")
    rows = await db.update("bookings", {"id": f"eq.{bid}"}, {"status": body.status})
    if not rows:
        raise HTTPException(404, "Topilmadi")
    return rows[0]


@app.get("/api/admin/report")
async def admin_report(period: str = "day", branch_id: int = 0, barber_id: int = 0, a: dict = Depends(staff_user)):
    now = datetime.now(TZ)
    start = now.replace(hour=0, minute=0, second=0, microsecond=0)
    if period == "month":
        start = start.replace(day=1)
        end = (start + timedelta(days=32)).replace(day=1)
    else:
        end = start + timedelta(days=1)
    params = [
        ("start_at", f"gte.{start.isoformat()}"), ("start_at", f"lt.{end.isoformat()}"),
        ("select", "status,total_price,svc_names,service_id,service_ids"),
    ]
    if branch_id:
        params.append(("branch_id", f"eq.{branch_id}"))
    if barber_id:
        params.append(("barber_id", f"eq.{barber_id}"))
    if a["scope"] is not None:
        params.append(("barber_id", in_list(a["scope"])))
    rows = await db.select("bookings", params)
    counts = {"booked": 0, "done": 0, "no_show": 0, "cancelled": 0}
    revenue, by_service, names_map = 0, {}, None
    for r in rows:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
        if r["status"] == "done":
            revenue += r["total_price"] or 0
            if names_map is None:
                names_map = {x["id"]: x["name"] for x in await db.select("services", {"select": "id,name"})}
            # nomlar xizmat ID'laridan olinadi (eski yozuvlarda svc_names bo'sh bo'lishi mumkin)
            ids = r.get("service_ids") or ([r["service_id"]] if r.get("service_id") else [])
            names = [names_map[i] for i in ids if i in names_map]
            if not names:
                names = [n.strip() for n in (r.get("svc_names") or "").split(", ") if n.strip()]
            for n in names:
                by_service[n] = by_service.get(n, 0) + 1
    top = sorted(by_service.items(), key=lambda x: -x[1])
    return {"period": period, "revenue": revenue, "counts": counts,
            "by_service": [{"name": n, "count": c} for n, c in top]}


# ---- analytics (hisobot): bitta hisob-kitob funksiyasi -- dashboard (keyin eksport va avtomatik hisobot) shundan foydalanadi
async def select_all(table, params, page=1000):
    """PostgREST 1000 qatordan ortig'ini qaytarmaydi, shuning uchun sahifalab o'qiymiz."""
    out, off = [], 0
    while True:
        rows = await db.select(table, list(params) + [("order", "id.asc"), ("limit", page), ("offset", off)])
        out += rows
        if len(rows) < page:
            return out
        off += page


def _pct(part, whole):
    return round(part * 100 / whole, 1) if whole else None


def compute_analytics(rows, d1, d2, now, prior_clients, svc_map, group_month):
    """Sof hisob-kitob (bazaga tegmaydi). rows -- davrdagi bronlar."""
    counts = {"pending": 0, "booked": 0, "done": 0, "no_show": 0, "cancelled": 0}
    elig = {"done": 0, "no_show": 0, "cancelled": 0}  # natijasi ma'lum, o'tib ketgan bronlar
    revenue, by_service, visits = 0, {}, {}
    if group_month:
        keys, y, m = [], d1.year, d1.month
        while (y, m) <= (d2.year, d2.month):
            keys.append(f"{y:04d}-{m:02d}")
            m += 1
            if m == 13:
                y, m = y + 1, 1
    else:
        keys = [(d1 + timedelta(days=i)).isoformat() for i in range((d2 - d1).days + 1)]
    series = {k: {"k": k, "revenue": 0, "done": 0, "cancelled": 0, "no_show": 0, "_c": set()} for k in keys}
    for r in rows:
        st = r["status"]
        counts[st] = counts.get(st, 0) + 1
        start = pdt(r["start_at"]).astimezone(TZ)
        b = series.get(start.strftime("%Y-%m") if group_month else start.date().isoformat())
        if start < now and st in elig:
            elig[st] += 1
        if st == "cancelled" and b:
            b["cancelled"] += 1
        if st == "no_show" and b:
            b["no_show"] += 1
        if st != "done":
            continue
        price = r.get("total_price") or 0
        revenue += price
        if b:
            b["revenue"] += price
            b["done"] += 1
            if r.get("client_id") is not None:
                b["_c"].add(r["client_id"])
        ids = r.get("service_ids") or ([r["service_id"]] if r.get("service_id") else [])
        names = [svc_map[i] for i in ids if i in svc_map]
        if not names:  # eski yozuvlar: matndan
            names = [n.strip() for n in (r.get("svc_names") or "").split(", ") if n.strip()]
        for n in names:
            by_service[n] = by_service.get(n, 0) + 1
        cid = r.get("client_id")
        if cid is not None:
            v = visits.setdefault(cid, [0, 0])
            v[0] += 1
            v[1] += price
    done_n = counts["done"]
    unique = len(visits)
    returning = sum(1 for c in visits if c in prior_clients)
    denom = sum(elig.values())
    vals = [{**{k: v for k, v in x.items() if k != "_c"}, "unique": len(x["_c"])} for x in series.values()]
    best = max(vals, key=lambda x: x["done"], default=None)
    worst = max(vals, key=lambda x: x["cancelled"], default=None)
    top_raw = sorted(((c, v[0], v[1]) for c, v in visits.items() if v[0] >= 2), key=lambda x: (-x[1], -x[2]))[:5]
    return {
        "revenue": revenue,
        "avg_check": round(revenue / done_n) if done_n else None,
        "counts": counts,
        "upcoming": counts["pending"] + counts["booked"],
        "customers": {"unique": unique, "returning": returning, "new": unique - returning,
                      "returning_rate": _pct(returning, unique)},
        "rates": {"eligible": denom, "completion": _pct(elig["done"], denom),
                  "cancelled": _pct(elig["cancelled"], denom), "no_show": _pct(elig["no_show"], denom)},
        "by_service": [{"name": n, "count": c} for n, c in sorted(by_service.items(), key=lambda x: -x[1])],
        "series": vals,
        "group": "month" if group_month else "day",
        "insights": {
            "best": {"k": best["k"], "count": best["done"]} if best and best["done"] else None,
            "worst_cancel": {"k": worst["k"], "count": worst["cancelled"]} if worst and worst["cancelled"] else None,
        },
        "top_raw": top_raw,
    }


def _parse_range(start, end, prev_start="", prev_end=""):
    try:
        d1, d2 = date.fromisoformat(start), date.fromisoformat(end)
        p1 = date.fromisoformat(prev_start) if prev_start and prev_end else None
        p2 = date.fromisoformat(prev_end) if prev_start and prev_end else None
    except ValueError:
        raise HTTPException(400, "Sana noto'g'ri")
    if d1 > d2 or (p1 and p1 > p2):
        raise HTTPException(400, "Boshlanish sanasi tugash sanasidan keyin bo'lishi mumkin emas")
    days = (d2 - d1).days + 1
    if days > 366:
        raise HTTPException(400, "Davr 1 yildan oshmasligi kerak")
    if not p1:  # oldingi teng davr
        p2 = d1 - timedelta(days=1)
        p1 = p2 - timedelta(days=days - 1)
    return d1, d2, p1, p2


async def build_report(d1, d2, p1, p2, branch_id=0, barber_id=0, scope=None, daily=False):
    """Dashboard va eksport uchun yagona hisob-kitob. Qaytaradi: (natija, bronlar qatorlari)."""
    days = (d2 - d1).days + 1
    flt = []
    if branch_id:
        flt.append(("branch_id", f"eq.{branch_id}"))
    if barber_id:
        flt.append(("barber_id", f"eq.{barber_id}"))
    if scope is not None:
        flt.append(("barber_id", in_list(scope)))

    lo, hi = datetime.combine(d1, dtime.min, TZ), datetime.combine(d2 + timedelta(days=1), dtime.min, TZ)
    plo, phi = datetime.combine(p1, dtime.min, TZ), datetime.combine(p2 + timedelta(days=1), dtime.min, TZ)
    rows = await select_all("bookings", [
        ("start_at", f"gte.{lo.isoformat()}"), ("start_at", f"lt.{hi.isoformat()}"),
        ("select", "status,total_price,start_at,client_id,service_id,service_ids,svc_names,barber_id,branch_id")] + flt)
    prev_rows = await select_all("bookings", [
        ("status", "eq.done"), ("start_at", f"gte.{plo.isoformat()}"), ("start_at", f"lt.{phi.isoformat()}"),
        ("select", "total_price")] + flt)
    prev_rev = sum(r["total_price"] or 0 for r in prev_rows)

    done_clients = sorted({r["client_id"] for r in rows if r["status"] == "done" and r.get("client_id") is not None})
    prior = set()  # shu davrdan oldin bajarilgan vizit bo'lgan mijozlar (do'kon bo'yicha)
    for i in range(0, len(done_clients), 100):
        pf = [("status", "eq.done"), ("start_at", f"lt.{lo.isoformat()}"),
              ("client_id", in_list(done_clients[i:i + 100])), ("select", "client_id")]
        if scope is not None:
            pf.append(("barber_id", in_list(scope)))
        prior |= {x["client_id"] for x in await select_all("bookings", pf)}

    svc_map = {}
    if any(r["status"] == "done" for r in rows):
        svc_map = {x["id"]: x["name"] for x in await db.select("services", {"select": "id,name"})}

    out = compute_analytics(rows, d1, d2, datetime.now(TZ), prior, svc_map, days > 62 and not daily)
    top_raw = out.pop("top_raw")
    names = {}
    if top_raw:
        us = await db.select("users", [("telegram_id", in_list([c for c, _, _ in top_raw])), ("select", "telegram_id,name")])
        names = {u["telegram_id"]: u["name"] for u in us}
    out["top_customers"] = [{"name": names.get(c) or "—", "visits": n, "revenue": rev} for c, n, rev in top_raw]
    diff = out["revenue"] - prev_rev
    out["compare"] = {"current": out["revenue"], "previous": prev_rev, "diff": diff,
                      "pct": _pct(diff, prev_rev) if prev_rev else None,
                      "prev_start": p1.isoformat(), "prev_end": p2.isoformat()}
    out["range"] = {"start": d1.isoformat(), "end": d2.isoformat(), "days": days}
    return out, rows


@app.get("/api/admin/analytics")
async def admin_analytics(start: str, end: str, prev_start: str = "", prev_end: str = "",
                          branch_id: int = 0, barber_id: int = 0, a: dict = Depends(staff_user)):
    d1, d2, p1, p2 = _parse_range(start, end, prev_start, prev_end)
    out, _ = await build_report(d1, d2, p1, p2, branch_id, barber_id, a["scope"])
    return out


# ---- eksport (Excel / PDF): dashboard bilan bir xil hisob-kitob, fayl botdan Telegramga yuboriladi
def compute_breakdowns(rows, barber_names, branch_names):
    """Ustalar va filiallar bo'yicha bajarilgan vizitlar (sof hisob-kitob)."""
    bb, br = {}, {}
    for r in rows:
        if r["status"] != "done":
            continue
        price = r.get("total_price") or 0
        x = bb.setdefault(r.get("barber_id"), [0, 0])
        x[0] += 1
        x[1] += price
        g = br.setdefault(r.get("branch_id"), {"done": 0, "revenue": 0, "c": set()})
        g["done"] += 1
        g["revenue"] += price
        if r.get("client_id") is not None:
            g["c"].add(r["client_id"])
    by_barber = sorted(({"name": barber_names.get(k) or "—", "done": v[0], "revenue": v[1],
                         "avg": round(v[1] / v[0]) if v[0] else None} for k, v in bb.items()),
                       key=lambda x: -x["revenue"])
    by_branch = sorted(({"name": branch_names.get(k) or "—", "done": v["done"], "revenue": v["revenue"],
                         "unique": len(v["c"])} for k, v in br.items()), key=lambda x: -x["revenue"])
    return by_barber, by_branch


def _dmy(iso):
    return f"{iso[8:10]}.{iso[5:7]}.{iso[:4]}"


def _money(n):
    return f"{int(n):,}".replace(",", " ") + " so‘m"


def _pc(p):
    return "—" if p is None else f"{p}%"


def make_xlsx(rep) -> bytes:
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    o, cmp_ = rep["out"], rep["out"]["compare"]
    SUM, PCT = '#,##0" so‘m"', '0.0"%"'
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="14223A")

    def table(ws, headers, data, fmts, row0=1):
        for j, hd in enumerate(headers, 1):
            c = ws.cell(row=row0, column=j, value=hd)
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, rw in enumerate(data, row0 + 1):
            for j, v in enumerate(rw, 1):
                c = ws.cell(row=i, column=j, value=v)
                if fmts[j - 1] and isinstance(v, (int, float)):
                    c.number_format = fmts[j - 1]
        for j, hd in enumerate(headers, 1):
            w = max([len(str(hd))] + [len(str(rw[j - 1])) for rw in data]) + 3
            ws.column_dimensions[get_column_letter(j)].width = min(max(w, 12), 40)
        ws.freeze_panes = ws.cell(row=row0 + 1, column=1)

    wb = Workbook()
    ws = wb.active
    ws.title = "Xulosa"
    ws["A1"] = "Hisobot"
    ws["A1"].font = Font(bold=True, size=16)
    cu, rt = o["customers"], o["rates"]
    items = [
        ("Davr", f"{_dmy(o['range']['start'])} – {_dmy(o['range']['end'])}", None),
        ("Tayyorlangan vaqt", rep["generated"], None),
        ("Filial", rep["branch_label"], None),
        ("Usta", rep["barber_label"], None),
        (None, None, None),
        ("Jami tushum", o["revenue"], SUM),
        ("Bajarilgan bronlar", o["counts"]["done"], "0"),
        ("Kutilayotgan bronlar (tushumga kirmaydi)", o["upcoming"], "0"),
        ("Bekor qilingan bronlar", o["counts"]["cancelled"], "0"),
        ("Kelmagan bronlar", o["counts"]["no_show"], "0"),
        ("O‘rtacha chek", o["avg_check"] if o["avg_check"] is not None else "—", SUM),
        (None, None, None),
        ("Kelgan mijozlar (noyob)", cu["unique"], "0"),
        ("Yangi mijozlar", cu["new"], "0"),
        ("Qaytgan mijozlar", cu["returning"], "0"),
        ("Qaytganlar ulushi", cu["returning_rate"] if cu["returning_rate"] is not None else "—", PCT),
        (None, None, None),
        ("Bajarilish foizi", rt["completion"] if rt["completion"] is not None else "—", PCT),
        ("Bekor qilish foizi", rt["cancelled"] if rt["cancelled"] is not None else "—", PCT),
        ("Kelmaslik foizi", rt["no_show"] if rt["no_show"] is not None else "—", PCT),
        ("Foizlar bazasi (o‘tgan, natijasi ma’lum bronlar)", rt["eligible"], "0"),
        (None, None, None),
        (f"Oldingi davr ({_dmy(cmp_['prev_start'])} – {_dmy(cmp_['prev_end'])}) tushumi", cmp_["previous"], SUM),
        ("Farq", cmp_["diff"], SUM),
        ("O‘zgarish", cmp_["pct"] if cmp_["pct"] is not None else "—", PCT),
    ]
    for i, (k, v, f) in enumerate(items, 3):
        if k is None:
            continue
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=v)
        c.alignment = Alignment(horizontal="left")
        if f and isinstance(v, (int, float)):
            c.number_format = f
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 30

    table(wb.create_sheet("Kunlar"), ["Sana", "Tushum", "Bajarilgan", "Bekor", "Kelmagan", "Noyob mijozlar"],
          [[_dmy(x["k"]), x["revenue"], x["done"], x["cancelled"], x["no_show"], x["unique"]] for x in o["series"]],
          [None, SUM, "0", "0", "0", "0"])
    table(wb.create_sheet("Ustalar"), ["Usta", "Bajarilgan", "Tushum", "O‘rtacha chek"],
          [[x["name"], x["done"], x["revenue"], x["avg"] if x["avg"] is not None else "—"] for x in rep["by_barber"]] or [["Ma’lumot yo‘q", "", "", ""]],
          [None, "0", SUM, SUM])
    table(wb.create_sheet("Xizmatlar"), ["Xizmat", "Bajarilgan soni"],
          [[x["name"], x["count"]] for x in o["by_service"]] or [["Ma’lumot yo‘q", ""]], [None, "0"])
    if rep["show_branches"]:
        table(wb.create_sheet("Filiallar"), ["Filial", "Bajarilgan", "Tushum", "Noyob mijozlar"],
              [[x["name"], x["done"], x["revenue"], x["unique"]] for x in rep["by_branch"]] or [["Ma’lumot yo‘q", "", "", ""]],
              [None, "0", SUM, "0"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


_FONT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fonts")


def _pdf_fonts():
    """Unicode shrift (o‘, g‘, kirill). Topilmasa -- standart shrift va noma'lum belgilar '?' bo'ladi."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for reg, bold in [(os.path.join(_FONT_DIR, "DejaVuSans.ttf"), os.path.join(_FONT_DIR, "DejaVuSans-Bold.ttf")),
                      ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
                      ("/usr/share/fonts/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf")]:
        if os.path.exists(reg) and os.path.exists(bold):
            if "RptSans" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("RptSans", reg))
                pdfmetrics.registerFont(TTFont("RptSans-B", bold))
            return "RptSans", "RptSans-B", True
    return "Helvetica", "Helvetica-Bold", False


def make_pdf(rep) -> bytes:
    from io import BytesIO
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font, bold, uni = _pdf_fonts()
    INK, BLUE, GREY = colors.HexColor("#14223A"), colors.HexColor("#1F4E9C"), colors.HexColor("#ECEFF3")

    def tx(v):
        v = str(v)
        return escape(v if uni else v.encode("cp1252", "replace").decode("cp1252"))

    st_n = ParagraphStyle("n", fontName=font, fontSize=9, leading=12, textColor=INK)
    st_b = ParagraphStyle("b", parent=st_n, fontName=bold)
    st_h = ParagraphStyle("h", fontName=bold, fontSize=12, leading=15, textColor=BLUE, spaceBefore=10, spaceAfter=4, keepWithNext=1)
    st_t = ParagraphStyle("t", fontName=bold, fontSize=20, leading=24, textColor=INK)
    st_m = ParagraphStyle("m", parent=st_n, textColor=colors.HexColor("#667489"))
    st_hd = ParagraphStyle("hd", parent=st_n, fontName=bold, textColor=colors.white)
    W = A4[0] - 30 * mm

    def tbl(headers, data, widths, num_cols=()):
        cell = lambda v, j, base: Paragraph(tx(v), ParagraphStyle("c", parent=base, alignment=2 if j in num_cols else 0))
        body = [[cell(h, j, st_hd) for j, h in enumerate(headers)]] + [[cell(v, j, st_n) for j, v in enumerate(r)] for r in data]
        t = Table(body, colWidths=[w * W for w in widths], repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), INK), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, GREY]),
                               ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#DDE3EA")),
                               ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        return t

    o, c = rep["out"], rep["out"]["compare"]
    cu, rt = o["customers"], o["rates"]
    el = [Paragraph(tx("Hisobot"), st_t),
          Paragraph(tx(f"{_dmy(o['range']['start'])} – {_dmy(o['range']['end'])}"), st_b),
          Paragraph(tx(f"Filial: {rep['branch_label']}   ·   Usta: {rep['barber_label']}"), st_m),
          Paragraph(tx(f"Tayyorlangan: {rep['generated']}"), st_m), Spacer(1, 6)]

    el += [Paragraph(tx("Asosiy ko‘rsatkichlar"), st_h), tbl(["Ko‘rsatkich", "Qiymat"], [
        ["Jami tushum", _money(o["revenue"])],
        ["O‘rtacha chek", "—" if o["avg_check"] is None else _money(o["avg_check"])],
        ["Bajarilgan bronlar", o["counts"]["done"]],
        ["Kutilayotgan bronlar (tushumga kirmaydi)", o["upcoming"]],
        ["Bekor qilingan bronlar", o["counts"]["cancelled"]],
        ["Kelmagan bronlar", o["counts"]["no_show"]]], [0.62, 0.38], (1,))]

    sign = "+" if c["diff"] > 0 else ("−" if c["diff"] < 0 else "")
    el += [Paragraph(tx("Oldingi davr bilan taqqoslash"), st_h), tbl(["", "Qiymat"], [
        [f"Oldingi davr ({_dmy(c['prev_start'])} – {_dmy(c['prev_end'])})", _money(c["previous"])],
        ["Farq", sign + _money(abs(c["diff"]))],
        ["O‘zgarish", "—" if c["pct"] is None else f"{'+' if c['pct'] > 0 else ('−' if c['pct'] < 0 else '')}{abs(c['pct'])}%"]],
        [0.62, 0.38], (1,))]

    el += [Paragraph(tx("Mijozlar"), st_h), tbl(["Ko‘rsatkich", "Qiymat"], [
        ["Kelgan mijozlar (noyob)", cu["unique"]], ["Yangi mijozlar", cu["new"]],
        ["Qaytgan mijozlar", cu["returning"]], ["Qaytganlar ulushi", _pc(cu["returning_rate"])]], [0.62, 0.38], (1,))]
    if o["top_customers"]:
        el += [Spacer(1, 4), tbl(["Eng ko‘p kelganlar", "Vizit", "Tushum"],
                                 [[x["name"], x["visits"], _money(x["revenue"])] for x in o["top_customers"]], [0.5, 0.15, 0.35], (1, 2))]

    el += [Paragraph(tx("Bronlar natijasi"), st_h), tbl(["Ko‘rsatkich", "Qiymat"], [
        ["Bajarilish foizi", _pc(rt["completion"])], ["Bekor qilish foizi", _pc(rt["cancelled"])],
        ["Kelmaslik foizi", _pc(rt["no_show"])]], [0.62, 0.38], (1,)),
        Paragraph(tx(f"Foizlar faqat vaqti o‘tgan va natijasi ma’lum bronlardan hisoblangan ({rt['eligible']} ta)."), st_m)]

    el += [Paragraph(tx("Kunlik statistika"), st_h), tbl(["Sana", "Tushum", "Bajarilgan", "Bekor", "Kelmagan", "Mijoz"],
           [[_dmy(x["k"]), _money(x["revenue"]), x["done"], x["cancelled"], x["no_show"], x["unique"]] for x in o["series"]],
           [0.17, 0.27, 0.16, 0.12, 0.15, 0.13], (1, 2, 3, 4, 5))]
    if rep["by_barber"]:
        el += [Paragraph(tx("Ustalar"), st_h), tbl(["Usta", "Bajarilgan", "Tushum", "O‘rtacha chek"],
               [[x["name"], x["done"], _money(x["revenue"]), "—" if x["avg"] is None else _money(x["avg"])] for x in rep["by_barber"]],
               [0.3, 0.16, 0.27, 0.27], (1, 2, 3))]
    if o["by_service"]:
        el += [Paragraph(tx("Xizmatlar"), st_h), tbl(["Xizmat", "Bajarilgan soni"],
               [[x["name"], x["count"]] for x in o["by_service"]], [0.7, 0.3], (1,))]
    if rep["show_branches"] and rep["by_branch"]:
        el += [Paragraph(tx("Filiallar"), st_h), tbl(["Filial", "Bajarilgan", "Tushum", "Mijoz"],
               [[x["name"], x["done"], _money(x["revenue"]), x["unique"]] for x in rep["by_branch"]],
               [0.34, 0.16, 0.32, 0.18], (1, 2, 3))]

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#667489"))
        canvas.drawRightString(A4[0] - 15 * mm, 10 * mm, tx(f"Sahifa {doc.page}"))
        canvas.restoreState()

    buf = BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=14 * mm, bottomMargin=16 * mm,
                      title="Hisobot").build(el, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


_export_last: dict = {}


@app.post("/api/admin/export")
async def admin_export(fmt: str, start: str, end: str, branch_id: int = 0, barber_id: int = 0, a: dict = Depends(admin_user)):
    if fmt not in ("xlsx", "pdf"):
        raise HTTPException(400, "Format noto'g'ri")
    d1, d2, p1, p2 = _parse_range(start, end)
    now_m = _time.monotonic()
    if now_m - _export_last.get(a["telegram_id"], -99.0) < 8:
        raise HTTPException(429, "Biroz kuting, oldingi fayl tayyorlanmoqda")
    _export_last[a["telegram_id"]] = now_m
    out, rows = await build_report(d1, d2, p1, p2, branch_id, barber_id, None, daily=True)
    branches = {x["id"]: x["name"] for x in await db.select("branches", {"select": "id,name"})}
    barbers = {x["id"]: x["name"] for x in await db.select("barbers", {"select": "id,name"})}
    by_barber, by_branch = compute_breakdowns(rows, barbers, branches)
    data = {
        "out": out, "by_barber": by_barber, "by_branch": by_branch, "show_branches": not branch_id,
        "branch_label": branches.get(branch_id, "—") if branch_id else "Hammasi",
        "barber_label": barbers.get(barber_id, "—") if barber_id else "Hammasi",
        "generated": datetime.now(TZ).strftime("%d.%m.%Y %H:%M"),
    }
    blob = await asyncio.to_thread(make_xlsx if fmt == "xlsx" else make_pdf, data)
    fname = f"hisobot_{d1.isoformat()}_{d2.isoformat()}.{fmt}"
    try:
        await bot.send_document(a["telegram_id"], BufferedInputFile(blob, filename=fname),
                                caption=f"Hisobot: {_dmy(d1.isoformat())} – {_dmy(d2.isoformat())}")
    except Exception:
        log.exception("eksport yuborilmadi")
        raise HTTPException(502, "Faylni Telegramga yuborib bo'lmadi. Botni ishga tushirganingizni tekshiring.")
    return {"ok": True}


# ---- services
@app.get("/api/admin/services")
async def admin_services(a: dict = Depends(staff_user)):
    params = {"order": "id.asc"}
    if a["scope"] is not None:
        params["barber_id"] = in_list(a["scope"])
    return await db.select("services", params)


@app.post("/api/admin/services")
async def admin_add_service(body: ServiceIn, a: dict = Depends(staff_user)):
    need_own_barber(a, body.barber_id)
    return (await db.insert("services", body.model_dump()))[0]


@app.patch("/api/admin/services/{sid}")
async def admin_edit_service(sid: int, body: ServicePatch, a: dict = Depends(staff_user)):
    data = body.model_dump(exclude_none=True)
    if not data:
        raise HTTPException(400, "O'zgartirish yo'q")
    if a["scope"] is not None:
        cur = await db.select("services", {"id": f"eq.{sid}", "select": "barber_id"})
        if not cur:
            raise HTTPException(404, "Topilmadi")
        need_own_barber(a, cur[0]["barber_id"])
    rows = await db.update("services", {"id": f"eq.{sid}"}, data)
    if not rows:
        raise HTTPException(404, "Topilmadi")
    return rows[0]


# ---- shop info
@app.get("/api/admin/shop")
async def admin_get_shop(a: dict = Depends(admin_user)):
    rows = await db.select("shop", {"id": "eq.1", "select": SHOP_COLS})
    return shop_out(rows[0])


@app.patch("/api/admin/shop")
async def admin_edit_shop(body: ShopPatch, a: dict = Depends(admin_user)):
    data = body.model_dump(exclude_unset=True)
    if "name" in data and not (data["name"] or "").strip():
        raise HTTPException(400, "Nomini kiriting")
    check_coords(data)
    if not data:
        raise HTTPException(400, "O'zgartirish yo'q")
    rows = await db.update("shop", {"id": "eq.1"}, data)
    return shop_out((await db.select("shop", {"id": "eq.1", "select": SHOP_COLS}))[0]) if rows else {}


@app.put("/api/admin/shop/logo")
async def admin_set_logo(body: ImageIn, a: dict = Depends(admin_user)):
    check_image(body.data)
    cur = await db.select("shop", {"id": "eq.1", "select": "logo_ver"})
    ver = (cur[0]["logo_ver"] if cur else 0) + 1
    await db.update("shop", {"id": "eq.1"}, {"logo_data": body.data, "logo_ver": ver})
    return {"logo_url": img_url("/api/img/logo", ver)}


# ---- branches
@app.get("/api/admin/branches")
async def admin_branches(a: dict = Depends(staff_user)):
    params = {"order": "id.asc", "select": BRANCH_COLS}
    if a["scope"] is not None:
        mine = await db.select("barbers", {"id": in_list(a["scope"]), "select": "branch_id"})
        bids = sorted({r["branch_id"] for r in mine if r.get("branch_id") is not None})
        if not bids:
            return []
        params["id"] = in_list(bids)
    rows = await db.select("branches", params)
    return [branch_out(r) for r in rows]


@app.post("/api/admin/branches")
async def admin_add_branch(body: BranchIn, a: dict = Depends(admin_user)):
    data = body.model_dump()
    if not data["name"].strip():
        raise HTTPException(400, "Nomini kiriting")
    check_coords(data)
    row = (await db.insert("branches", data))[0]
    hours = [{"branch_id": row["id"], "weekday": w, "start_time": "09:00", "end_time": "20:00", "is_off": False}
             for w in range(6)]
    hours.append({"branch_id": row["id"], "weekday": 6, "start_time": None, "end_time": None, "is_off": True})
    await db.insert("work_hours", hours)
    return branch_out(row)


@app.patch("/api/admin/branches/{bid}")
async def admin_edit_branch(bid: int, body: BranchPatch, a: dict = Depends(admin_user)):
    data = body.model_dump(exclude_unset=True)
    if "name" in data and not (data["name"] or "").strip():
        raise HTTPException(400, "Nomini kiriting")
    if data.get("active") is None:
        data.pop("active", None)
    check_coords(data)
    if not data:
        raise HTTPException(400, "O'zgartirish yo'q")
    rows = await db.update("branches", {"id": f"eq.{bid}"}, data)
    if not rows:
        raise HTTPException(404, "Topilmadi")
    return branch_out(rows[0])


@app.put("/api/admin/branches/{bid}/photo")
async def admin_branch_photo(bid: int, body: ImageIn, a: dict = Depends(admin_user)):
    check_image(body.data)
    cur = await db.select("branches", {"id": f"eq.{bid}", "select": "photo_ver"})
    if not cur:
        raise HTTPException(404, "Topilmadi")
    ver = (cur[0]["photo_ver"] or 0) + 1
    await db.update("branches", {"id": f"eq.{bid}"}, {"photo_data": body.data, "photo_ver": ver})
    return {"photo_url": img_url(f"/api/img/branch/{bid}", ver)}


# ---- work hours (per branch)
@app.get("/api/admin/work-hours")
async def admin_hours(branch_id: int, a: dict = Depends(admin_user)):
    return await db.select("work_hours", {"branch_id": f"eq.{branch_id}", "order": "weekday.asc"})


@app.put("/api/admin/work-hours/{branch_id}/{weekday}")
async def admin_set_hours(branch_id: int, weekday: int, body: HoursIn, a: dict = Depends(admin_user)):
    if not 0 <= weekday <= 6:
        raise HTTPException(400, "Kun noto'g'ri")
    if not body.is_off and not (body.start_time and body.end_time):
        raise HTTPException(400, "Boshlanish va tugash vaqtini kiriting")
    data = ({"start_time": None, "end_time": None, "is_off": True} if body.is_off
            else {"start_time": body.start_time, "end_time": body.end_time, "is_off": False})
    rows = await db.update("work_hours", {"branch_id": f"eq.{branch_id}", "weekday": f"eq.{weekday}"}, data)
    if not rows:
        rows = await db.insert("work_hours", {"branch_id": branch_id, "weekday": weekday, **data})
    return rows[0]


# ---- barbers admin + days off
@app.get("/api/admin/barbers")
async def admin_barbers(a: dict = Depends(staff_user)):
    params = {"order": "id.asc", "select": BARBER_COLS_ADM}
    if a["scope"] is not None:
        params["id"] = in_list(a["scope"])
    return [barber_admin_out(r) for r in await db.select("barbers", params)]


@app.post("/api/admin/barbers")
async def admin_add_barber(body: BarberIn, a: dict = Depends(admin_user)):
    if not body.name.strip():
        raise HTTPException(400, "Ismni kiriting")
    return barber_admin_out((await db.insert("barbers", {"branch_id": body.branch_id, "name": body.name.strip(), "name_ru": (body.name_ru or "").strip() or None}))[0])


@app.patch("/api/admin/barbers/{bid}")
async def admin_edit_barber(bid: int, body: BarberPatch, a: dict = Depends(staff_user)):
    need_own_barber(a, bid)
    data = body.model_dump(exclude_none=True)
    if a["scope"] is not None:  # usta faqat ism va ruscha ismini o'zgartira oladi
        data = {k: v for k, v in data.items() if k in ("name", "name_ru")}
        if "name" in data and not data["name"].strip():
            raise HTTPException(400, "Ismni kiriting")
    if not data:
        raise HTTPException(400, "O'zgartirish yo'q")
    rows = await db.update("barbers", {"id": f"eq.{bid}"}, data)
    if not rows:
        raise HTTPException(404, "Topilmadi")
    return barber_admin_out(rows[0])


@app.post("/api/admin/barbers/{bid}/invite")
async def barber_invite(bid: int, a: dict = Depends(admin_user)):
    if not await db.select("barbers", {"id": f"eq.{bid}", "select": "id"}):
        raise HTTPException(404, "Topilmadi")
    code = secrets.token_urlsafe(8)
    await db.insert("invites", {"code": code, "barber_id": bid})
    return {"link": f"https://t.me/{BOT_USERNAME}?start=brb_{code}"}


@app.delete("/api/admin/barbers/{bid}/telegram")
async def barber_unlink(bid: int, a: dict = Depends(admin_user)):
    rows = await db.update("barbers", {"id": f"eq.{bid}"}, {"telegram_id": None})
    if not rows:
        raise HTTPException(404, "Topilmadi")
    return {"ok": True}


@app.put("/api/admin/barbers/{bid}/photo")
async def admin_barber_photo(bid: int, body: ImageIn, a: dict = Depends(staff_user)):
    need_own_barber(a, bid)
    check_image(body.data)
    cur = await db.select("barbers", {"id": f"eq.{bid}", "select": "photo_ver"})
    if not cur:
        raise HTTPException(404, "Topilmadi")
    ver = (cur[0]["photo_ver"] or 0) + 1
    await db.update("barbers", {"id": f"eq.{bid}"}, {"photo_data": body.data, "photo_ver": ver})
    return {"photo_url": img_url(f"/api/img/barber/{bid}", ver)}


@app.get("/api/admin/days-off")
async def admin_days_off(barber_id: int, a: dict = Depends(staff_user)):
    need_own_barber(a, barber_id)
    return await db.select("days_off", {"barber_id": f"eq.{barber_id}", "order": "day.asc"})


@app.post("/api/admin/days-off")
async def admin_set_day_off(body: DayOffIn, a: dict = Depends(staff_user)):
    need_own_barber(a, body.barber_id)
    try:
        day = date.fromisoformat(body.day).isoformat()
    except ValueError:
        raise HTTPException(400, "Sana noto'g'ri")
    if body.off:
        try:
            await db.insert("days_off", {"barber_id": body.barber_id, "day": day})
        except DBError as e:
            if "23505" not in e.body:
                raise
    else:
        await db.delete("days_off", {"barber_id": f"eq.{body.barber_id}", "day": f"eq.{day}"})
    return {"ok": True}


# ---- owner only: admins management
@app.post("/api/admin/invite")
async def create_invite(o: dict = Depends(owner_user)):
    code = secrets.token_urlsafe(8)
    await db.insert("invites", {"code": code})
    return {"link": f"https://t.me/{BOT_USERNAME}?start=adm_{code}"}


@app.get("/api/admin/admins")
async def list_admins(o: dict = Depends(owner_user)):
    return await db.select("users", {"role": "eq.admin", "select": "telegram_id,name,phone"})


@app.delete("/api/admin/admins/{tid}")
async def remove_admin(tid: int, o: dict = Depends(owner_user)):
    rows = await db.update("users", {"telegram_id": f"eq.{tid}", "role": "eq.admin"}, {"role": "client"})
    if not rows:
        raise HTTPException(404, "Admin topilmadi")
    return {"ok": True}
